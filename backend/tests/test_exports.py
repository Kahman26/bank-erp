import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.db import SessionLocal
from app.models import Role
from app.services.atm_pdf import FONT_CANDIDATES
from seed.generator import Generator
from seed.writer import ensure_base, id_allocator, write_db


def _seed(count: int = 40) -> None:
    with SessionLocal() as db:
        ref = ensure_base(db)
        write_db(db, Generator(ref, id_allocator(db), 31).run(count))
        db.commit()


def _sheet(content: bytes):
    wb = load_workbook(io.BytesIO(content))
    return wb, list(wb.worksheets[0].iter_rows(values_only=True))


def test_atm_export_respects_filters_and_reimports_cleanly(login_as) -> None:
    _seed()
    admin = login_as(Role.ADMIN)
    ref = admin.get("/api/v1/reference").json()
    in_service = next(d["id"] for d in ref["dicts"]["atm_status"] if d["code"] == "in_service")

    r = admin.get("/api/v1/atms/export.xlsx", params={"status_id": in_service})
    assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxml")
    wb, rows = _sheet(r.content)
    expected = admin.get("/api/v1/atms", params={"status_id": in_service}).json()["total"]
    assert len(rows) - 1 == expected
    assert rows[0][0] == "Код банкомата" and "Сервисный вендор" in rows[0]
    assert all(row[rows[0].index("Статус")] == "В эксплуатации" for row in rows[1:])
    assert wb["Параметры"]["B2"].value == str(expected)

    # выгрузка без фильтров загружается обратно без ошибок и ничего не меняет
    full = admin.get("/api/v1/atms/export.xlsx").content
    files = {"file": ("atms.xlsx", full)}
    preview = admin.post("/api/v1/import/atms/preview", files=files).json()
    assert preview["errors"] == 0, preview["rows"][:3]
    assert preview["update"] == 40
    result = admin.post("/api/v1/import/atms/commit", files={"file": ("atms.xlsx", full)}).json()
    assert result["unchanged"] == 40, result


def test_certificates_export(login_as) -> None:
    _seed(10)
    c = login_as(Role.VIEWER)
    r = c.get("/api/v1/certificates/export.xlsx", params={"status": "expired"})
    _, rows = _sheet(r.content)
    assert rows[0][0] == "Объект"
    assert all(row[9] == "Истёк" for row in rows[1:])
    assert len(rows) - 1 == c.get("/api/v1/certificates", params={"status": "expired"}).json()["total"]


@pytest.mark.skipif(not any(Path(f[0]).is_file() for f in FONT_CANDIDATES), reason="нет шрифта с кириллицей")
def test_card_pdf(login_as) -> None:
    _seed(5)
    c = login_as(Role.VIEWER)
    atm_id = c.get("/api/v1/atms/map").json()["points"][0][0]
    r = c.get(f"/api/v1/atms/{atm_id}/card.pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")
    assert len(r.content) > 5000
    assert c.get("/api/v1/atms/999999/card.pdf").status_code == 404


def test_dashboard_is_consistent(login_as) -> None:
    _seed(60)
    c = login_as(Role.VIEWER)
    d = c.get("/api/v1/reports/dashboard").json()
    t = d["totals"]
    assert t["atms"] == 60
    assert sum(s["count"] for s in d["by_status"]) == 60
    assert sum(r["total"] for r in d["by_region"]) == t["in_network"]
    assert sum(m["total"] for m in d["by_model"]) == t["in_network"]
    assert sum(r["in_repair"] for r in d["by_region"]) == t["in_repair"]
    assert d["repairs_by_month"] and all(len(m["month"]) == 7 for m in d["repairs_by_month"])
    assert len(d["repairs_by_month"]) <= 12
    assert any(v["performer"] == "Инженер банка" for v in d["vendors"]) or d["vendors"]


def test_report_exports(login_as) -> None:
    _seed(20)
    c = login_as(Role.VIEWER)
    for name in ("by-region", "by-model", "vendors", "repairs-by-month", "end-of-support", "warranty"):
        r = c.get(f"/api/v1/reports/{name}.xlsx")
        assert r.status_code == 200, name
        _sheet(r.content)
    assert c.get("/api/v1/reports/nope.xlsx").status_code == 404
