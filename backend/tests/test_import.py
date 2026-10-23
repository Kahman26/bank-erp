import csv
import io

from openpyxl import load_workbook
from sqlalchemy import func, select

from app.db import SessionLocal
from app.models import Atm, AtmStatusHistory, Placement, Role
from seed.generator import Generator
from seed.writer import dirty_rows, ensure_base, id_allocator


def _rows(count: int = 30, seed: int = 5) -> list[dict[str, str]]:
    """Строки CSV от генератора — без записи в БД (база для импорта пустая)."""
    with SessionLocal() as db:
        ref = ensure_base(db)
        db.commit()  # модели, филиалы и вендоры нужны импорту
        return Generator(ref, id_allocator(db), seed).run(count).csv


def _csv(rows: list[dict[str, str]], delimiter: str = ",") -> bytes:
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(rows[0]), delimiter=delimiter)
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def _upload(client, path: str, content: bytes, name: str = "atms.csv", **form):
    return client.post(f"/api/v1/import/atms/{path}", files={"file": (name, content)}, data=form)


def test_clean_file_imports_and_is_idempotent(login_as) -> None:
    rows = _rows()
    admin = login_as(Role.ADMIN)
    data = _csv(rows, delimiter=";")  # разделитель определяется сам

    preview = _upload(admin, "preview", data).json()
    assert preview["total"] == 30 and preview["create"] == 30 and preview["errors"] == 0

    r = _upload(admin, "commit", data)
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 30

    with SessionLocal() as db:
        assert db.scalar(select(func.count(Atm.id))) == 30
        placed = sum(1 for row in rows if row["latitude"])
        assert db.scalar(select(func.count(Placement.id)).where(Placement.placed_to.is_(None))) == placed
        history = db.scalar(select(func.count(AtmStatusHistory.id)))

    # повторная загрузка того же файла ничего не меняет
    again = _upload(admin, "commit", data).json()
    assert again == {"created": 0, "updated": 0, "unchanged": 30, "skipped": 0, "moved": 0}
    with SessionLocal() as db:
        assert db.scalar(select(func.count(AtmStatusHistory.id))) == history


def test_moved_atm_gets_new_placement(login_as) -> None:
    rows = _rows(5)
    admin = login_as(Role.ADMIN)
    _upload(admin, "commit", _csv(rows))
    row = next(r for r in rows if r["latitude"])
    row["latitude"] = f"{float(row['latitude']) + 0.01:.6f}"
    row["address"] = "ул. Новая, 1"
    result = _upload(admin, "commit", _csv(rows)).json()
    assert result["moved"] == 1 and result["updated"] == 1
    with SessionLocal() as db:
        atm_id = db.scalar(select(Atm.id).where(Atm.atm_code == row["atm_code"]))
        placements = db.scalars(select(Placement).where(Placement.atm_id == atm_id)).all()
    assert len(placements) == 2 and sum(p.placed_to is None for p in placements) == 1


def test_dirty_file_reports_errors_per_row(login_as) -> None:
    rows = dirty_rows(_rows(80), seed=5)
    expected = {r["atm_code"]: r.pop("_error") for r in rows}
    admin = login_as(Role.ADMIN)
    preview = _upload(admin, "preview", _csv(rows)).json()

    by_note = {}
    for r in preview["rows"]:
        by_note.setdefault(r["atm_code"], []).append(r)
    fixable = {
        "регион без «г.» — должен распознаться",
        "дата в формате ДД.ММ.ГГГГ — должна распознаться",
        "запятая в координате — должна распознаться",
    }
    for code, note in expected.items():
        if not note:
            continue
        reported = by_note.get(code, [])
        assert reported, f"не замечено: {note}"
        if note in fixable:
            assert any(r["action"] != "error" and r["warnings"] for r in reported), note
        else:
            assert any(r["action"] == "error" for r in reported), note
    assert preview["errors"] >= 9

    # без пропуска ошибок ничего не загружается
    assert _upload(admin, "commit", _csv(rows)).status_code == 400
    with SessionLocal() as db:
        assert db.scalar(select(func.count(Atm.id))) == 0
    result = _upload(admin, "commit", _csv(rows), skip_errors="true").json()
    assert result["skipped"] == preview["errors"] and result["created"] == preview["create"]


def test_bad_files(login_as) -> None:
    admin = login_as(Role.ADMIN)
    assert _upload(admin, "preview", b"a,b\n1,2\n").status_code == 400  # нет обязательных колонок
    assert _upload(admin, "preview", b"x", name="atms.pdf").status_code == 400
    assert _upload(login_as(Role.OPERATOR), "preview", b"x").status_code == 403


def test_template_and_xlsx_roundtrip(login_as) -> None:
    rows = _rows(3)
    admin = login_as(Role.ADMIN)
    r = admin.get("/api/v1/import/atms/template")
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Банкоматы", "Описание"]

    # файл Excel с русскими заголовками тоже читается
    ws = wb["Банкоматы"]
    ws.delete_rows(1, ws.max_row)
    labels = {row[0].value: row[1].value for row in wb["Описание"].iter_rows(min_row=2)}
    ws.append([labels[c] for c in rows[0]])
    for row in rows:
        ws.append(list(row.values()))
    buf = io.BytesIO()
    wb.save(buf)
    preview = _upload(admin, "preview", buf.getvalue(), name="atms.xlsx").json()
    assert preview["total"] == 3 and preview["errors"] == 0, preview
