"""Единое оформление Excel-выгрузок: шапка, фильтр, закреплённая строка, лист с параметрами."""

import io
from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(bold=True, color="FFFFFF")
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _cell(value: Any) -> Any:
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, bool):
        return "да" if value else "нет"
    return value


def build_xlsx(
    title: str,
    columns: Sequence[tuple[str, int]],
    rows: Iterable[Sequence[Any]],
    params: Sequence[tuple[str, str]] = (),
) -> bytes:
    """columns — (заголовок, ширина); params — что выгружено и кем, на отдельный лист."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([c[0] for c in columns])
    for i, (_, width) in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    count = 0
    for row in rows:
        ws.append([_cell(v) for v in row])
        count += 1
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, date | datetime):
                cell.number_format = "DD.MM.YYYY HH:MM" if isinstance(cell.value, datetime) else "DD.MM.YYYY"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{count + 1}"

    info = wb.create_sheet("Параметры")
    for row in [("Выгрузка", title), ("Строк", str(count)), *params]:
        info.append(list(row))
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 80
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def content_disposition(filename: str) -> dict[str, str]:
    # имя файла латиницей: кириллица в заголовках требует RFC 5987 и ломается в старых браузерах
    return {"Content-Disposition": f'attachment; filename="{filename}"'}
