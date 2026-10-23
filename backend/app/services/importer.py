"""Загрузка банкоматов из CSV/XLSX (формат — раздел 5 набора полей).

Проверка и применение разделены: интерфейс сначала показывает результат проверки по строкам,
затем отправляет тот же файл на применение — и он проверяется заново. Повторная загрузка того же
файла ничего не меняет: запись ищется по коду банкомата.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from app.models import (
    Atm,
    AtmModel,
    AtmProcessing,
    AtmStatusHistory,
    Branch,
    Currency,
    DictItem,
    Location,
    Manufacturer,
    PaymentScheme,
    Placement,
    Processing,
    Region,
)

MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
UZ_LAT, UZ_LON = (37.0, 46.0), (55.0, 74.0)

# колонка → (подпись, обязательна)
COLUMNS: dict[str, tuple[str, bool]] = {
    "atm_code": ("Код банкомата", True),
    "terminal_id_own": ("Terminal ID — собственный процессинг", False),
    "terminal_id_uzcard": ("Terminal ID — Uzcard", False),
    "terminal_id_humo": ("Terminal ID — Humo", False),
    "merchant_id_own": ("Merchant ID — собственный процессинг", False),
    "merchant_id_uzcard": ("Merchant ID — Uzcard", False),
    "merchant_id_humo": ("Merchant ID — Humo", False),
    "serial_number": ("Серийный номер", True),
    "manufacturer": ("Производитель", True),
    "model": ("Модель", True),
    "device_type": ("Тип устройства", True),
    "install_type": ("Формат установки", True),
    "status": ("Статус", True),
    "purchase_date": ("Дата покупки", False),
    "commissioning_date": ("Дата ввода", False),
    "branch": ("Филиал", True),
    "region": ("Регион", True),
    "city_district": ("Город / район", False),
    "address": ("Адрес", False),
    "landmark": ("Ориентир", False),
    "latitude": ("Широта", False),
    "longitude": ("Долгота", False),
    "place_type": ("Тип места", False),
    "working_hours": ("Режим работы", False),
    "currencies": ("Валюты", False),
    "accepted_schemes": ("Принимаемые карты", False),
    "screen_size_in": ("Диагональ экрана", False),
    "screen_resolution": ("Разрешение", False),
    "cassette_count": ("Кол-во кассет", False),
    "cash_in": ("Приём наличных", False),
    "contactless": ("Бесконтактный ридер", False),
}
LABEL_TO_COLUMN = {label.lower(): col for col, (label, _) in COLUMNS.items()}
PROCESSING_CODES = ("own", "uzcard", "humo")
YES = {"да", "yes", "y", "true", "1", "+", "есть"}
NO = {"нет", "no", "n", "false", "0", "-", ""}


class ImportFileError(ValueError):
    """Файл целиком не читается: формат, размер, нет нужных колонок."""


# --- чтение файла ------------------------------------------------------------------------------------------


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def read_table(filename: str, content: bytes) -> list[dict[str, str]]:
    if len(content) > MAX_BYTES:
        raise ImportFileError("Файл больше 5 МБ")
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
        except Exception as e:  # повреждённый или не тот формат
            raise ImportFileError("Не удалось прочитать Excel-файл") from e
        raw = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    elif name.endswith((".csv", ".txt")):
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ImportFileError("Неизвестная кодировка файла — сохраните CSV в UTF-8")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        raw = [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ImportFileError("Поддерживаются файлы .csv и .xlsx")

    raw = [r for r in raw if any(r)]
    if not raw:
        raise ImportFileError("Файл пустой")
    header = [h.strip().lower() for h in raw[0]]
    columns = [h if h in COLUMNS else LABEL_TO_COLUMN.get(h, h) for h in header]
    missing = [COLUMNS[c][0] for c, (_, req) in COLUMNS.items() if req and c not in columns]
    if missing:
        raise ImportFileError("В файле нет обязательных колонок: " + ", ".join(missing))
    if len(raw) - 1 > MAX_ROWS:
        raise ImportFileError(f"Слишком много строк: больше {MAX_ROWS}")
    return [{col: (r[i] if i < len(r) else "") for i, col in enumerate(columns) if col in COLUMNS} for r in raw[1:]]


# --- проверка ------------------------------------------------------------------------------------------------


@dataclass
class RowResult:
    line: int  # номер строки в файле (с учётом заголовка)
    atm_code: str
    action: str = "create"  # create / update / error
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    values: dict[str, Any] = field(default_factory=dict)


def _norm(s: str) -> str:
    s = s.lower().replace("ё", "е")
    s = re.sub(r"\b(г|обл|область|респ|республика)\b\.?", " ", s)
    return re.sub(r"[^\w]+", " ", s).strip()


class Lookup:
    """Поиск значения справочника по названию, коду или тексту в скобках, без учёта регистра."""

    def __init__(self, items: list[tuple[Any, list[str]]]) -> None:
        self.exact: dict[str, Any] = {}
        self.loose: dict[str, list[Any]] = {}
        for value, keys in items:
            for k in keys:
                if k:
                    self.exact.setdefault(k.strip().lower(), value)
                    self.loose.setdefault(_norm(k), []).append(value)

    def find(self, raw: str) -> tuple[Any | None, bool]:
        """(значение, найдено ли только после нормализации)."""
        key = raw.strip().lower()
        if key in self.exact:
            return self.exact[key], False
        cands = list(dict.fromkeys(self.loose.get(_norm(raw), [])))
        return (cands[0], True) if len(cands) == 1 else (None, False)


def _dict_lookup(items: list[DictItem]) -> Lookup:
    entries = []
    for d in items:
        keys = [d.name, d.code]
        if m := re.search(r"\(([^)]+)\)", d.name):
            keys.append(m.group(1))
        entries.append((d.id, keys))
    return Lookup(entries)


class Importer:
    def __init__(self, db: Session) -> None:
        self.db = db
        dicts = db.scalars(select(DictItem).where(DictItem.is_active)).all()
        by_kind: dict[str, list[DictItem]] = {}
        for d in dicts:
            by_kind.setdefault(d.kind, []).append(d)
        self.status = _dict_lookup(by_kind.get("atm_status", []))
        self.device_type = _dict_lookup(by_kind.get("device_type", []))
        self.install_format = _dict_lookup(by_kind.get("install_format", []))
        self.place_type = _dict_lookup(by_kind.get("place_type", []))
        self.regions = Lookup([(r.id, [r.name, r.code]) for r in db.scalars(select(Region))])
        self.branches = Lookup([(b.id, [b.name]) for b in db.scalars(select(Branch).where(Branch.is_active))])
        self.currencies = {c.code.upper() for c in db.scalars(select(Currency))}
        self.schemes = Lookup([(s.id, [s.name, s.code]) for s in db.scalars(select(PaymentScheme))])
        self.processings = {p.code: p.id for p in db.scalars(select(Processing))}
        self.models = {
            (mf.lower(), m.model_name.lower()): m.id
            for m, mf in db.execute(select(AtmModel, Manufacturer.name).join(Manufacturer))
        }
        self.existing = {
            a.atm_code.lower(): a
            for a in db.scalars(select(Atm).options(selectinload(Atm.processings), selectinload(Atm.currencies)))
        }
        self.serials = {a.serial_number.lower(): a.atm_code.lower() for a in self.existing.values()}
        # (processing_id, terminal_id) → atm_code — для поиска дублей с уже существующими банкоматами
        self.terminals: dict[tuple[int, str], str] = {}
        self.merchants: dict[tuple[int, str], str] = {}
        for code, atm in self.existing.items():
            for p in atm.processings:
                self.terminals[(p.processing_id, p.terminal_id.lower())] = code
                self.merchants[(p.processing_id, p.merchant_id.lower())] = code

    # --- разбор отдельных значений -------------------------------------------------------------------------

    @staticmethod
    def _date(raw: str, label: str, res: RowResult) -> date | None:
        if not raw:
            return None
        for fmt, warn in (("%Y-%m-%d", False), ("%d.%m.%Y", True), ("%d/%m/%Y", True)):
            try:
                value = datetime.strptime(raw, fmt).date()
            except ValueError:
                continue
            if warn:
                res.warnings.append(f"{label}: дата «{raw}» распознана как {value:%d.%m.%Y}")
            return value
        res.errors.append(f"{label}: «{raw}» — не дата (ожидается ГГГГ-ММ-ДД)")
        return None

    @staticmethod
    def _decimal(raw: str, label: str, res: RowResult) -> Decimal | None:
        if not raw:
            return None
        text = raw.replace(" ", "")
        if "," in text and "." not in text:
            text = text.replace(",", ".")
            res.warnings.append(f"{label}: запятая заменена на точку")
        try:
            return Decimal(text)
        except InvalidOperation:
            res.errors.append(f"{label}: «{raw}» — не число")
            return None

    @staticmethod
    def _bool(raw: str, label: str, res: RowResult) -> bool | None:
        key = raw.strip().lower()
        if key in YES:
            return True
        if key in NO:
            return False if key else None
        res.errors.append(f"{label}: «{raw}» — ожидается «да» или «нет»")
        return None

    def _pick(self, lookup: Lookup, raw: str, label: str, res: RowResult, required: bool = True) -> Any | None:
        if not raw:
            if required:
                res.errors.append(f"{label}: не заполнено")
            return None
        value, fuzzy = lookup.find(raw)
        if value is None:
            res.errors.append(f"{label}: «{raw}» нет в справочнике")
        elif fuzzy:
            res.warnings.append(f"{label}: «{raw}» сопоставлено со справочником")
        return value

    # --- проверка строки ----------------------------------------------------------------------------------------

    def validate(self, rows: list[dict[str, str]]) -> list[RowResult]:
        results = []
        seen_codes: dict[str, int] = {}
        seen_serials: dict[str, int] = {}
        seen_terminals: dict[tuple[int, str], int] = {}
        for i, row in enumerate(rows, start=2):
            res = RowResult(line=i, atm_code=row.get("atm_code", ""))
            self._validate_row(row, res)
            code = res.atm_code.lower()
            if code:
                if code in seen_codes:
                    res.errors.append(f"Код банкомата повторяется (строка {seen_codes[code]})")
                seen_codes.setdefault(code, i)
            serial = row.get("serial_number", "").lower()
            if serial:
                if serial in seen_serials:
                    res.errors.append(f"Серийный номер повторяется (строка {seen_serials[serial]})")
                seen_serials.setdefault(serial, i)
                owner = self.serials.get(serial)
                if owner and owner != code:
                    res.errors.append(f"Серийный номер уже у банкомата {owner.upper()}")
            for proc_id, tid, _ in res.values.get("processings", []):
                key = (proc_id, tid.lower())
                if key in seen_terminals:
                    res.errors.append(f"Terminal ID {tid} повторяется в процессинге (строка {seen_terminals[key]})")
                seen_terminals.setdefault(key, i)
            if res.errors:
                res.action = "error"
            elif code in self.existing:
                res.action = "update"
            results.append(res)
        return results

    def _validate_row(self, row: dict[str, str], res: RowResult) -> None:
        v = res.values
        g = row.get
        if not res.atm_code:
            res.errors.append("Код банкомата: не заполнено")
        if not g("serial_number"):
            res.errors.append("Серийный номер: не заполнено")
        v["serial_number"] = g("serial_number", "")

        model_id = self.models.get((g("manufacturer", "").strip().lower(), g("model", "").strip().lower()))
        if g("model") and model_id is None:
            res.errors.append(
                f"Модель «{g('manufacturer', '')} {g('model', '')}» не найдена — заведите её в разделе «Модели»"
            )
        v["model_id"] = model_id
        v["device_type_id"] = self._pick(self.device_type, g("device_type", ""), "Тип устройства", res)
        v["install_format_id"] = self._pick(self.install_format, g("install_type", ""), "Формат установки", res)
        v["status_id"] = self._pick(self.status, g("status", ""), "Статус", res)
        v["branch_id"] = self._pick(self.branches, g("branch", ""), "Филиал", res)
        v["region_id"] = self._pick(self.regions, g("region", ""), "Регион", res)
        v["purchase_date"] = self._date(g("purchase_date", ""), "Дата покупки", res)
        v["commissioning_date"] = self._date(g("commissioning_date", ""), "Дата ввода", res)
        v["screen_size_in"] = self._decimal(g("screen_size_in", ""), "Диагональ экрана", res)
        v["screen_resolution"] = (g("screen_resolution") or "").replace("×", "x") or None
        count = self._decimal(g("cassette_count", ""), "Кол-во кассет", res)
        v["cassette_count"] = int(count) if count is not None else None
        v["contactless"] = self._bool(g("contactless", ""), "Бесконтактный ридер", res)

        currencies = [c.strip().upper() for c in re.split(r"[;,]", g("currencies", "")) if c.strip()]
        unknown = [c for c in currencies if c not in self.currencies]
        if unknown:
            res.errors.append(f"Валюты: неизвестно {', '.join(unknown)}")
        v["currencies"] = currencies or None

        schemes = []
        for s in (x for x in re.split(r"[;,]", g("accepted_schemes", "")) if x.strip()):
            sid = self._pick(self.schemes, s, "Принимаемые карты", res)
            if sid:
                schemes.append(sid)
        v["scheme_ids"] = schemes or None

        processings = []
        for code in PROCESSING_CODES:
            tid, mid = g(f"terminal_id_{code}", ""), g(f"merchant_id_{code}", "")
            if not tid and not mid:
                continue
            if not tid or not mid:
                res.errors.append(f"Процессинг {code}: нужны и Terminal ID, и Merchant ID")
                continue
            pid = self.processings[code]
            for registry, value, label in ((self.terminals, tid, "Terminal ID"), (self.merchants, mid, "Merchant ID")):
                owner = registry.get((pid, value.lower()))
                if owner and owner != res.atm_code.lower():
                    res.errors.append(f"{label} {value} в процессинге {code} уже у банкомата {owner.upper()}")
            processings.append((pid, tid, mid))
        v["processings"] = processings

        lat = self._decimal(g("latitude", ""), "Широта", res)
        lon = self._decimal(g("longitude", ""), "Долгота", res)
        if (lat is None) != (lon is None):
            res.errors.append("Координаты: нужны и широта, и долгота")
        elif lat is not None and lon is not None:
            if not (UZ_LAT[0] <= lat <= UZ_LAT[1] and UZ_LON[0] <= lon <= UZ_LON[1]):
                swapped = UZ_LAT[0] <= lon <= UZ_LAT[1] and UZ_LON[0] <= lat <= UZ_LON[1]
                res.errors.append(
                    "Координаты вне Узбекистана" + (" — похоже, перепутаны широта и долгота" if swapped else "")
                )
            for col, label in (("address", "Адрес"), ("city_district", "Город / район")):
                if not g(col):
                    res.errors.append(f"{label}: не заполнено (обязательно, если указаны координаты)")
            v["location"] = {
                "latitude": lat,
                "longitude": lon,
                "address": g("address", ""),
                "city_district": g("city_district", ""),
                "landmark": g("landmark") or None,
                "working_hours": g("working_hours") or "—",
                "place_type_id": self._pick(self.place_type, g("place_type", ""), "Тип места", res),
            }

    # --- применение ------------------------------------------------------------------------------------------------

    def apply(self, results: list[RowResult], user_id: int) -> dict[str, int]:
        today = date.today()
        stats = {"created": 0, "updated": 0, "unchanged": 0, "skipped": 0, "moved": 0}
        currencies = {c.code: c for c in self.db.scalars(select(Currency))}
        schemes = {s.id: s for s in self.db.scalars(select(PaymentScheme))}
        for res in results:
            if res.action == "error":
                stats["skipped"] += 1
                continue
            v = res.values
            atm = self.existing.get(res.atm_code.lower())
            created = atm is None
            if created:
                atm = Atm(atm_code=res.atm_code, status_id=v["status_id"])
                self.db.add(atm)
            changed = created
            for f in (
                "model_id",
                "device_type_id",
                "install_format_id",
                "branch_id",
                "purchase_date",
                "commissioning_date",
                "screen_size_in",
                "screen_resolution",
                "cassette_count",
            ):
                if v[f] is not None and getattr(atm, f) != v[f]:
                    setattr(atm, f, v[f])
                    changed = True
            if v["contactless"] is not None and atm.contactless != v["contactless"]:
                atm.contactless = v["contactless"]
                changed = True
            if atm.serial_number != v["serial_number"]:
                atm.serial_number = v["serial_number"]
                changed = True
            if v["currencies"] is not None and sorted(c.code for c in atm.currencies) != sorted(v["currencies"]):
                atm.currencies = [currencies[c] for c in v["currencies"]]
                changed = True
            if v["scheme_ids"] is not None and sorted(s.id for s in atm.schemes) != sorted(v["scheme_ids"]):
                atm.schemes = [schemes[s] for s in v["scheme_ids"]]
                changed = True
            self.db.flush()

            if created or atm.status_id != v["status_id"]:
                atm.status_id = v["status_id"]
                self.db.add(
                    AtmStatusHistory(
                        atm_id=atm.id, status_id=v["status_id"], changed_by_id=user_id, comment="Загрузка из файла"
                    )
                )
                changed = True

            current = {p.processing_id: p for p in atm.processings} if not created else {}
            for pid, tid, mid in v["processings"]:
                proc = current.get(pid)
                if proc is None:
                    self.db.add(AtmProcessing(atm_id=atm.id, processing_id=pid, terminal_id=tid, merchant_id=mid))
                    changed = True
                elif (proc.terminal_id, proc.merchant_id) != (tid, mid):
                    proc.terminal_id, proc.merchant_id = tid, mid
                    changed = True

            if loc := v.get("location"):
                moved, loc_changed = self._place(atm, loc, v["region_id"], v["commissioning_date"] or today, user_id)
                stats["moved"] += moved
                changed = changed or moved or loc_changed

            stats["created" if created else "updated" if changed else "unchanged"] += 1
        return stats

    def _place(self, atm: Atm, loc: dict, region_id: int, start: date, user_id: int) -> tuple[bool, bool]:
        """Ставит банкомат на место из файла. (перемещён ли, изменились ли данные текущего места)."""
        current = self.db.scalar(select(Placement).where(Placement.atm_id == atm.id, Placement.placed_to.is_(None)))
        fields = {**loc, "region_id": region_id}
        if current is not None:
            location = self.db.get(Location, current.location_id)
            same_place = abs(location.latitude - loc["latitude"]) < Decimal("0.0001") and abs(
                location.longitude - loc["longitude"]
            ) < Decimal("0.0001")
            if same_place:
                changed = False
                for f, value in fields.items():
                    if value is not None and getattr(location, f) != value:
                        setattr(location, f, value)
                        changed = True
                return False, changed
            today = date.today()
            current.placed_to = max(today, current.placed_from)
            start = current.placed_to
            self.db.flush()
        new_loc = Location(
            name=fields["address"],
            indoor=True,
            is_24h=fields["working_hours"].lower().startswith(("кругл", "24")),
            **{k: v for k, v in fields.items() if v is not None},
        )
        if new_loc.place_type_id is None:
            new_loc.place_type_id = self.db.scalar(
                select(DictItem.id).where(DictItem.kind == "place_type", DictItem.code == "street")
            )
        self.db.add(new_loc)
        self.db.flush()
        self.db.add(Placement(atm_id=atm.id, location_id=new_loc.id, placed_from=start, created_by_id=user_id))
        return current is not None, False


def template_xlsx(db: Session) -> bytes:
    """Шаблон: лист с колонками и примером строки + лист с описанием и допустимыми значениями."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Банкоматы"
    ws.append(list(COLUMNS))
    ws.append(
        [
            "ATM-TAS-0154",
            "10010154",
            "UZ050154",
            "7710154H",
            "900154",
            "9010154",
            "8810154",
            "HY1234567",
            "Hyosung TNS",
            "MX8300D",
            "Ресайклер",
            "В помещении (lobby)",
            "В эксплуатации",
            "2024-03-15",
            "2024-04-02",
            "Юнусабадский филиал",
            "г. Ташкент",
            "Юнусабадский район",
            "ул. Амира Темура, 1",
            "Напротив станции метро",
            41.311081,
            69.240562,
            "ТЦ",
            "Круглосуточно",
            "UZS;USD",
            "Uzcard;Humo;Visa;Mastercard;UnionPay",
            19,
            "1280x1024",
            4,
            "да",
            "да",
        ]
    )
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(14, len(str(col[0].value)) + 2)

    info = wb.create_sheet("Описание")
    info.append(["Колонка", "Поле", "Обязательно", "Допустимые значения"])
    values: dict[str, str] = {}
    for kind, column in (
        ("device_type", "device_type"),
        ("install_format", "install_type"),
        ("atm_status", "status"),
        ("place_type", "place_type"),
    ):
        names = db.scalars(
            select(DictItem.name).where(DictItem.kind == kind, DictItem.is_active).order_by(DictItem.sort)
        ).all()
        values[column] = "; ".join(names)
    values["region"] = "; ".join(db.scalars(select(Region.name).order_by(Region.sort)))
    values["branch"] = "название филиала из справочника"
    values["manufacturer"] = "; ".join(db.scalars(select(Manufacturer.name).order_by(Manufacturer.name)))
    values["model"] = "модель из раздела «Модели»"
    values["currencies"] = "; ".join(db.scalars(select(Currency.code).order_by(Currency.sort))) + " — через «;»"
    values["accepted_schemes"] = "; ".join(db.scalars(select(PaymentScheme.name).order_by(PaymentScheme.sort)))
    values["purchase_date"] = values["commissioning_date"] = "ГГГГ-ММ-ДД или ДД.ММ.ГГГГ"
    values["latitude"] = "41.311081 — вместе с долготой, адресом и городом/районом"
    values["cash_in"] = values["contactless"] = "да / нет"
    for col, (label, required) in COLUMNS.items():
        info.append([col, label, "да" if required else "", values.get(col, "")])
    for letter, width in zip("ABCD", (22, 38, 12, 90), strict=True):
        info.column_dimensions[letter].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
